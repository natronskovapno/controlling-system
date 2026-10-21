"""
bb_parser.py — Bruto bilans parser (PDF + XLSX) for the Finansijski izveštaji module.

parse_file(path) -> dict:
  period_from, period_to (ISO), company, pib, fmt,
  lines: [{konto, naziv, ps_d, ps_p, tp_d, tp_p, uk_d, uk_p}]   (analytic kontos only)
  checks: [{name, ok, detail, args}], warnings: [(key, args)]

  Diagnostics are returned as a catalogue key plus its arguments rather than
  a finished sentence, so the upload preview can render them in either
  language. Older snapshots stored a plain string, which the templates still
  accept.
Saldo is not stored from the file; it is derived as uk_d - uk_p.
"""
import re
from datetime import date

NUM_RE = re.compile(r'-?\d{1,3}(?:\.\d{3})*,\d{2}')
KONTO_RE = re.compile(r'^\s*(\d{1,6})\b')
PERIOD_RE = re.compile(r'(\d{2})\.(\d{2})\.(\d{2,4})\s*-\s*(\d{2})\.(\d{2})\.(\d{2,4})')
CLOSING_KONTOS = ('599', '699')          # plus whole class 7


def _num(s):
    return float(s.replace('.', '').replace(',', '.'))


def _iso(d, m, y):
    y = int(y)
    if y < 100:
        y += 2000
    return date(y, int(m), int(d)).isoformat()


def _header_meta(text):
    meta = {'period_from': None, 'period_to': None, 'company': None, 'pib': None}
    flat = re.sub(r'\s+', ' ', text)
    m = PERIOD_RE.search(flat)
    if m:
        meta['period_from'] = _iso(*m.group(1, 2, 3))
        meta['period_to'] = _iso(*m.group(4, 5, 6))
    p = re.search(r'PIB:\s*(\d{9})', flat)
    if p:
        meta['pib'] = p.group(1)
    c = re.search(r'([A-ZČĆŠĐŽ][A-ZČĆŠĐŽ ]+DOO)', flat)
    if c:
        meta['company'] = re.sub(r'(?<=[ČĆŠĐŽ]) (?=[A-Z])', '', c.group(1)).strip()
    return meta


# ── PDF ──────────────────────────────────────────────────────────────
def _pdf_text(path):
    try:
        import pypdf
        r = pypdf.PdfReader(path)
        return '\n'.join(p.extract_text(extraction_mode='layout') or '' for p in r.pages)
    except ImportError:
        pass
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return '\n'.join(p.extract_text(layout=True) or '' for p in pdf.pages)
    except ImportError:
        raise RuntimeError('Za PDF je potreban paket pypdf (pip install pypdf) ili pdfplumber.')


def _parse_pdf(path):
    text = _pdf_text(path)
    rows, subtotals, classes, ukupno = [], {}, {}, None
    for line in text.splitlines():
        nums = NUM_RE.findall(line)
        if line.strip().upper().startswith('UKUPNO') and len(nums) >= 6:
            ukupno = [_num(x) for x in nums[:8]]
            continue
        km = KONTO_RE.match(line)
        if not km or len(nums) < 6:
            continue
        konto = km.group(1)
        first_num = NUM_RE.search(line)
        naziv = line[km.end():first_num.start()].strip()
        vals = [_num(x) for x in nums]
        if len(konto) == 1 and len(vals) == 8:
            classes[konto] = vals
        elif len(konto) == 3 and len(vals) == 8 and not naziv:
            subtotals[konto] = vals
        elif len(vals) in (6, 7):
            rows.append({'konto': konto, 'naziv': naziv or None,
                         'ps_d': vals[0], 'ps_p': vals[1], 'tp_d': vals[2],
                         'tp_p': vals[3], 'uk_d': vals[4], 'uk_p': vals[5]})
    return text[:3000], rows, subtotals, classes, ukupno


# ── XLSX ─────────────────────────────────────────────────────────────
def _parse_xlsx(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).active
    head = ' '.join(''.join(str(c) for c in r if c is not None)
                    for r in ws.iter_rows(min_row=1, max_row=10, values_only=True))
    rows, subtotals, classes, ukupno = [], {}, {}, None
    for r in ws.iter_rows(min_row=1, values_only=True):
        k = r[0]
        if k is None:
            continue
        k = str(k).strip()
        vals = [float(v) if isinstance(v, (int, float)) else 0.0 for v in r[2:10]]
        if k.upper().startswith('UKUPNO'):
            ukupno = vals
            continue
        if not k.isdigit():
            continue
        naziv = (str(r[1]).strip() if r[1] else '') or None
        # xlsx has no blank-description marker for subtotals; a 3-digit konto whose
        # 4-digit children precede it is a subtotal — resolved after the loop.
        rows.append({'konto': k, 'naziv': naziv, 'ps_d': vals[0], 'ps_p': vals[1],
                     'tp_d': vals[2], 'tp_p': vals[3], 'uk_d': vals[4], 'uk_p': vals[5]})
    out = []
    for i, row in enumerate(rows):
        k = row['konto']
        if len(k) == 1:
            classes[k] = [row[c] for c in ('ps_d', 'ps_p', 'tp_d', 'tp_p', 'uk_d', 'uk_p')]
            continue
        if len(k) == 3:
            has_children = any(o['konto'].startswith(k) and len(o['konto']) > 3 for o in out)
            dup_prev = out and out[-1]['konto'] == k
            if has_children or dup_prev:
                subtotals[k] = [row[c] for c in ('ps_d', 'ps_p', 'tp_d', 'tp_p', 'uk_d', 'uk_p')]
                continue
        out.append(row)
    return head, out, subtotals, classes, ukupno


# ── validation ───────────────────────────────────────────────────────
def _sr(x):
    return f'{x:,.2f}'.replace(',', ' ').replace('.', ',').replace(' ', '.')


def _close(a, b, tol=1.0):
    return abs(a - b) <= tol


def parse_file(path):
    ext = path.lower().rsplit('.', 1)[-1]
    if ext == 'pdf':
        head, rows, subtotals, classes, ukupno = _parse_pdf(path)
        fmt = 'PDF'
    elif ext in ('xlsx', 'xlsm'):
        head, rows, subtotals, classes, ukupno = _parse_xlsx(path)
        fmt = 'XLSX'
    else:
        raise RuntimeError('Podržani formati: PDF, XLSX.')

    meta = _header_meta(head)
    checks, warnings = [], []
    cols = ('ps_d', 'ps_p', 'tp_d', 'tp_p', 'uk_d', 'uk_p')

    checks.append({'name': 'Period prepoznat', 'ok': bool(meta['period_to']),
                   'detail': f"{meta['period_from']} → {meta['period_to']}"})
    checks.append({'name': 'Analitička konta', 'ok': len(rows) > 20,
                   'detail': '{a0} redova', 'args': {'a0': len(rows)}})

    if ukupno:
        tot = [sum(r[c] for r in rows) for c in cols]
        ok = all(_close(tot[i], ukupno[i], 5) for i in range(6))
        checks.append({'name': 'Zbir analitike = UKUPNO', 'ok': ok,
                       'detail': 'ukupan promet D {a0} / UKUPNO {a1}',
                       'args': {'a0': _sr(tot[4]), 'a1': _sr(ukupno[4])}})
        checks.append({'name': 'UKUPNO duguje = potražuje', 'ok': _close(ukupno[4], ukupno[5], 5),
                       'detail': f'{_sr(ukupno[4])} / {_sr(ukupno[5])}'})
    else:
        checks.append({'name': 'UKUPNO red pronađen', 'ok': False, 'detail': 'nije pronađen'})

    bad_classes = []
    for k, v in classes.items():
        s = [sum(r[c] for r in rows if r['konto'].startswith(k)) for c in cols]
        if not all(_close(s[i], v[i], 5) for i in range(6)):
            bad_classes.append(k)
    if bad_classes:
        class_detail = {'detail': 'odstupanje: klase {a0}', 'args': {'a0': ', '.join(bad_classes)}}
    else:
        class_detail = {'detail': 'sve klase se slažu'}
    checks.append({'name': 'Zbir po klasama', 'ok': not bad_classes and bool(classes),
                   **class_detail})

    bad_sub = []
    for k, v in subtotals.items():
        s = [sum(r[c] for r in rows if r['konto'][:3] == k) for c in cols]
        if not all(_close(s[i], v[i], 5) for i in range(6)):
            bad_sub.append(k)
    if bad_sub:
        warnings.append(('Sintetička konta sa odstupanjem: {a0}', {'a0': ', '.join(bad_sub[:15])}))

    is_closing = any(r['konto'].startswith('7') or r['konto'][:3] in CLOSING_KONTOS for r in rows)
    if is_closing:
        warnings.append(('Zaključni bilans: klase 5 i 6 su zatvorene — rezultat se čita iz prometa.', {}))
    no_name = sum(1 for r in rows if not r['naziv'])
    if no_name:
        warnings.append(('{a0} konta bez naziva u fajlu (naziv se preuzima iz drugih uvezenih bilansa).',
                         {'a0': no_name}))

    return {**meta, 'fmt': fmt, 'lines': rows, 'checks': checks, 'warnings': warnings,
            'is_closing': is_closing, 'ok': all(c['ok'] for c in checks)}
