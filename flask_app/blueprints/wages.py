"""
blueprints/wages.py
===================
Wages module — Garden Stage Controlling System

Landing view: records grouped by month (most recent first)
Secondary:    upload Excel templates (bookkeeper A/B, casual-labour C) with preview before commit

Schema: wages_daily
  id, date, employee_name, worker_type, venue, net_amount,
  tax_bracket, gross_cost, source_file, notes, created_at, fitid
"""

from flask import Blueprint, render_template, request, redirect, url_for, flash
from db import get_db
from datetime import date as _date
import os, hashlib
from i18n import tr

bp = Blueprint("wages", __name__)

# ── helpers ──────────────────────────────────────────────────────────

def _money(val):
    """Format number with dot thousands separator."""
    try:
        return f"{int(round(float(val))):,}".replace(",", ".")
    except (TypeError, ValueError):
        return "—"


def _get_monthly_summary(db, filters):
    """
    Returns rows grouped by year-month, with totals.
    Optionally filtered by employee, venue, worker_type.
    """
    where = "WHERE 1=1"
    params = []

    if filters.get("employee"):
        where += " AND employee_name = ?"
        params.append(filters["employee"])
    if filters.get("venue"):
        where += " AND venue = ?"
        params.append(filters["venue"])
    if filters.get("worker_type"):
        where += " AND worker_type = ?"
        params.append(filters["worker_type"])
    if filters.get("date_from"):
        where += " AND date >= ?"
        params.append(filters["date_from"])
    if filters.get("date_to"):
        where += " AND date <= ?"
        params.append(filters["date_to"])

    rows = db.execute(f"""
        SELECT
            strftime('%Y-%m', date)          AS month_key,
            strftime('%m/%Y', date)          AS month_label,
            COUNT(DISTINCT employee_name)    AS employee_count,
            COUNT(*)                         AS row_count,
            ROUND(SUM(net_amount), 0)        AS total_net,
            ROUND(SUM(gross_cost), 0)        AS total_gross
        FROM wages_daily
        {where}
        GROUP BY month_key
        ORDER BY month_key DESC
    """, params).fetchall()
    return [dict(r) for r in rows]


def _get_month_detail(db, month_key, filters):
    """Individual rows for a given YYYY-MM month."""
    where = "WHERE strftime('%Y-%m', date) = ?"
    params = [month_key]

    if filters.get("employee"):
        where += " AND employee_name = ?"
        params.append(filters["employee"])
    if filters.get("venue"):
        where += " AND venue = ?"
        params.append(filters["venue"])

    rows = db.execute(f"""
        SELECT id, date, employee_name, worker_type, venue,
               COALESCE(net_amount, 0)  AS net_amount,
               tax_bracket,
               COALESCE(gross_cost, 0) AS gross_cost,
               source_file, notes
        FROM wages_daily
        {where}
        ORDER BY date, employee_name
    """, params).fetchall()
    return [dict(r) for r in rows]


def _get_month_aggregated(db, month_key, filters):
    """Per-employee aggregate for a given YYYY-MM month."""
    where = "WHERE strftime('%Y-%m', date) = ?"
    params = [month_key]

    if filters.get("employee"):
        where += " AND employee_name = ?"
        params.append(filters["employee"])
    if filters.get("venue"):
        where += " AND venue = ?"
        params.append(filters["venue"])

    rows = db.execute(f"""
        SELECT employee_name, worker_type,
               COUNT(*)                      AS stavki,
               ROUND(SUM(net_amount), 0)     AS total_net,
               ROUND(SUM(gross_cost), 0)     AS total_gross,
               MIN(tax_bracket)              AS tax_bracket
        FROM wages_daily
        {where}
        GROUP BY employee_name, worker_type
        ORDER BY employee_name
    """, params).fetchall()
    return [dict(r) for r in rows]


def _ensure_indexes(db):
    """Create indexes on wages_daily if they don't exist yet (one-time, fast after first run)."""
    db.execute("CREATE INDEX IF NOT EXISTS idx_wages_date ON wages_daily(date)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_wages_emp ON wages_daily(employee_name)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_wages_date_emp ON wages_daily(date, employee_name, net_amount, gross_cost)")
    # No commit needed — CREATE INDEX is auto-committed by SQLite


def _get_filter_options(db):
    """Distinct values for filter dropdowns — single query instead of three."""
    rows = db.execute("""
        SELECT DISTINCT employee_name, venue, worker_type FROM wages_daily
        WHERE employee_name IS NOT NULL
        ORDER BY employee_name
    """).fetchall()
    seen_emp, seen_venue, seen_wtype = set(), set(), set()
    employees, venues, worker_types = [], [], []
    for r in rows:
        if r[0] and r[0] not in seen_emp:
            seen_emp.add(r[0]); employees.append(r[0])
        if r[1] and r[1] not in seen_venue:
            seen_venue.add(r[1]); venues.append(r[1])
        if r[2] and r[2] not in seen_wtype:
            seen_wtype.add(r[2]); worker_types.append(r[2])
    venues.sort(); worker_types.sort()
    return employees, venues, worker_types


def _kpis_from_monthly(monthly):
    """Derive KPI totals from already-fetched monthly summary rows — no extra DB query."""
    if not monthly:
        return {}
    return {
        "employees":   max((m["employee_count"] for m in monthly), default=0),
        "records":     sum(m["row_count"]   for m in monthly),
        "total_net":   sum(m["total_net"]   for m in monthly),
        "total_gross": sum(m["total_gross"] for m in monthly),
        "date_min":    min(m["month_key"]   for m in monthly),
        "date_max":    max(m["month_key"]   for m in monthly),
    }


def _aggregated_from_rows(expanded_rows):
    """
    Build per-employee aggregates from already-fetched expanded_rows in Python.
    Eliminates a second DB round-trip on Google Drive.
    """
    from collections import defaultdict
    buckets = defaultdict(lambda: {"stavki": 0, "total_net": 0.0, "total_gross": 0.0, "tax_bracket": None})
    for r in expanded_rows:
        key = (r["employee_name"], r["worker_type"] or "")
        b = buckets[key]
        b["stavki"]     += 1
        b["total_net"]  += float(r["net_amount"]  or 0)
        b["total_gross"] += float(r["gross_cost"] or 0)
        if b["tax_bracket"] is None:
            b["tax_bracket"] = r.get("tax_bracket")
    result = []
    for (emp, wtype), b in sorted(buckets.items(), key=lambda x: x[0][0]):
        result.append({
            "employee_name": emp,
            "worker_type":   wtype or None,
            "stavki":        b["stavki"],
            "total_net":     round(b["total_net"],  0),
            "total_gross":   round(b["total_gross"], 0),
            "tax_bracket":   b["tax_bracket"],
        })
    return result



def _get_kpis(db, filters):
    where = "WHERE 1=1"
    params = []
    if filters.get("employee"):
        where += " AND employee_name = ?"; params.append(filters["employee"])
    if filters.get("venue"):
        where += " AND venue = ?"; params.append(filters["venue"])
    if filters.get("worker_type"):
        where += " AND worker_type = ?"; params.append(filters["worker_type"])
    if filters.get("date_from"):
        where += " AND date >= ?"; params.append(filters["date_from"])
    if filters.get("date_to"):
        where += " AND date <= ?"; params.append(filters["date_to"])

    row = db.execute(f"""
        SELECT
            COUNT(DISTINCT employee_name)  AS employees,
            COUNT(*)                       AS records,
            ROUND(SUM(net_amount), 0)      AS total_net,
            ROUND(SUM(gross_cost), 0)      AS total_gross,
            MIN(date)                      AS date_min,
            MAX(date)                      AS date_max
        FROM wages_daily {where}
    """, params).fetchone()
    return dict(row) if row else {}


# ── Excel parsing ─────────────────────────────────────────────────────


# ── File parsing (openpyxl-based, matches 07b_ingest_wages_historical.py) ──────

def _parse_date(value):
    import re
    from datetime import date as _d, datetime as _dt
    if value is None:
        return None
    if isinstance(value, (_d, _dt)):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, float):
        value = f"{value:.10g}."
    s = str(value).strip()
    for pat in [r'^(\d{1,2})[./](\d{1,2})[./](\d{4})$', r'^(\d{4})-(\d{2})-(\d{2})$']:
        m = re.match(pat, s)
        if m:
            try:
                g = m.groups()
                if len(g[2]) == 4:
                    d = _dt(int(g[2]), int(g[1]), int(g[0]))
                else:
                    d = _dt(int(g[0]), int(g[1]), int(g[2]))
                return d.strftime('%Y-%m-%d')
            except (ValueError, IndexError):
                continue
    # Handle double-dot like "09.01..2025" → strip extra dots first
    s_clean = re.sub(r'\.{2,}', '.', s)
    m = re.search(r'(\d{1,2})\.(\d{1,2})\.(\d{3,4})', s_clean)
    if m:
        try:
            day, mon, yr = int(m.group(1)), int(m.group(2)), int(m.group(3))
            if yr < 100: yr += 2000
            return _dt(yr, mon, day).strftime('%Y-%m-%d')
        except ValueError:
            pass
    m2 = re.match(r'^(\d{1,2})\.(\d{1,2})\.$', s)
    if m2:
        return f"PARTIAL-{m2.group(2).zfill(2)}-{m2.group(1).zfill(2)}"
    return None


def _resolve_partial(partial, year):
    from datetime import datetime as _dt
    if not partial or not partial.startswith('PARTIAL-'):
        return partial
    try:
        _, mm, dd = partial.split('-')
        return _dt(year, int(mm), int(dd)).strftime('%Y-%m-%d')
    except ValueError:
        return None


def _get_bracket_for_date(db, employee_name, date_str):
    """Look up the applicable tax bracket from employee_bracket_history by date."""
    row = db.execute("""
        SELECT bracket FROM employee_bracket_history
        WHERE employee_name = ? AND valid_from <= ?
        ORDER BY valid_from DESC LIMIT 1
    """, [employee_name, date_str]).fetchone()
    return row["bracket"] if row else None


def _apply_tax(rows, db):
    """
    Add tax_bracket and gross_cost to each parsed row.
    Looks up bracket from employee_bracket_history by date.
    Bracket is a percentage string (e.g. '29', '76').
    gross_cost = net_amount * (1 + pct/100)
    """
    for r in rows:
        bracket = _get_bracket_for_date(db, r["employee_name"], r["date"])
        if bracket:
            try:
                pct = float(str(bracket).replace("%", "").strip())
                r["tax_bracket"] = str(bracket)
                r["gross_cost"]  = round(r["net_amount"] * (1 + pct / 100), 2)
            except (ValueError, TypeError):
                r["tax_bracket"] = "unknown"
                r["gross_cost"]  = 0
        else:
            r["tax_bracket"] = "unknown"
            r["gross_cost"]  = 0
        fitid_str = f"{r['date']}-{r['employee_name']}-{r['net_amount']}-{r['worker_type']}"
        r["fitid"] = hashlib.md5(fitid_str.encode()).hexdigest()[:16]
    return rows


def _detect_format_ws(ws):
    import re
    for r in range(1, min(ws.max_row+1, 15)):
        row_str = ' '.join(str(ws.cell(r, c).value or '') for c in range(1, ws.max_column+1))
        if any(kw in row_str for kw in ['Datumi', 'Cena sata']):
            return 'C'
        if re.search(r'Za\s+u[lp]?latu', row_str):
            return 'C'
    row2 = [ws.cell(2, c).value for c in range(1, min(ws.max_column+1, 25))]
    row2_str = ' '.join(str(v) for v in row2 if v)
    if re.search(r'(Ponedeljak|Utorak|Sreda|Četvrtak|Petak|Subota|Nedeljak)', row2_str):
        return 'B'
    if 'UKUPNO' in row2_str.upper():
        return 'B'
    row1 = [ws.cell(1, c).value for c in range(1, min(ws.max_column+1, 20))]
    names = sum(1 for v in row1[1:] if v and isinstance(v, str) and len(v.strip()) > 3)
    if names >= 2:
        return 'A'
    return 'A'


def _parse_format_a(ws, source_file):
    """Format A: names as column headers (row 1), dates as rows."""
    names = {}
    for c in range(2, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v and isinstance(v, str) and v.strip():
            names[c] = v.strip()
    rows = []
    for r in range(2, ws.max_row + 1):
        date_cell = ws.cell(r, 1).value
        if str(date_cell or '').strip().lower() in ('total', 'ukupno', 'porezi i doprinosi',
                                                     'ukupni trosak', 'ukupni trošak'):
            break
        parsed = _parse_date(date_cell)
        if not parsed or 'PARTIAL' in parsed:
            continue
        for col, name in names.items():
            val = ws.cell(r, col).value
            if val and isinstance(val, (int, float)) and float(val) > 0:
                rows.append({'date': parsed, 'employee_name': name,
                             'worker_type': 'konobar', 'venue': '',
                             'net_amount': float(val), 'source_file': source_file})
    return rows


def _parse_format_b(ws, source_file):
    """Format B: dates as column headers (row 2), names as rows.
    All non-date numeric columns are attributed to the nearest preceding date column.
    This handles second-shift amounts that land in unpaired columns.
    """
    import re
    date_cols = {}
    for c in range(4, ws.max_column + 1):
        v = ws.cell(2, c).value
        if v:
            d = _parse_date(v)
            if d and 'PARTIAL' not in d:
                date_cols[c] = d

    # Map every column to the nearest preceding date column
    col_to_date = {}
    last_date = None
    for c in range(4, ws.max_column + 1):
        if c in date_cols:
            last_date = date_cols[c]
        if last_date:
            col_to_date[c] = last_date

    # UKUPNO column = first col after last date col that has 'UKUPNO' header
    ukupno_col = None
    for c in range(4, ws.max_column + 1):
        v = str(ws.cell(2, c).value or '').upper()
        if 'UKUPNO' in v:
            ukupno_col = c
            break

    rows = []
    for r in range(5, ws.max_row + 1):
        name_cell = ws.cell(r, 2).value
        if not name_cell or not isinstance(name_cell, str) or not name_cell.strip():
            continue
        name = re.sub(r'[\ufeff\n]', '', name_cell).strip()
        if not name:
            continue
        for col, date_str in col_to_date.items():
            if ukupno_col and col >= ukupno_col:
                continue
            val = ws.cell(r, col).value
            if val and isinstance(val, (int, float)) and float(val) > 0:
                rows.append({'date': date_str, 'employee_name': name,
                             'worker_type': 'konobar', 'venue': '',
                             'net_amount': float(val), 'source_file': source_file})
    return rows


def _parse_format_c(ws, source_file):
    """Format C: casual labour — find header row with Datumi/Za uplatu."""
    import re
    from datetime import date as _d
    header_row = col_date = col_name = col_amount = None
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = str(ws.cell(r, c).value or '').strip()
            if v in ('Datumi', 'Datum'):
                header_row = r; col_date = c
            elif v in ('Imena', 'Ime'):
                col_name = c
            elif re.match(r'Za\s+u\S+atu', v, re.IGNORECASE) or v in ('Za isplatu', 'Iznos', 'Cena', 'Iznos', 'Ukupno'):
                col_amount = c
        if header_row:
            break
    if not header_row or not col_amount:
        return []
    rows = []
    current_date = None
    import re as _re
    m = _re.search(r'(\d{4})', source_file or '')
    year = int(m.group(1)) if m else _d.today().year
    for r in range(header_row + 1, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column+1)]
        if all(v is None for v in row_vals):
            continue
        if col_date:
            d_val = ws.cell(r, col_date).value
            if d_val:
                parsed = _parse_date(d_val)
                if parsed and 'PARTIAL' in parsed:
                    parsed = _resolve_partial(parsed, year)
                if parsed:
                    current_date = parsed
        amt_val = ws.cell(r, col_amount).value
        if not amt_val or not isinstance(amt_val, (int, float)) or float(amt_val) <= 0:
            continue
        name = None
        if col_name:
            n = ws.cell(r, col_name).value
            if n:
                name = str(n).strip().rstrip(',').strip()
        # Skip rows with no name — these are totals/summary rows
        if current_date and name:
            rows.append({'date': current_date, 'employee_name': name,
                         'worker_type': 'fizikalac', 'venue': '',
                         'net_amount': float(amt_val), 'source_file': source_file})
    return rows


def _parse_file(filepath, source_file, db):
    """Parse uploaded Excel file using openpyxl, detect format, return rows."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active
    except Exception as e:
        return [], tr("Greška pri čitanju fajla:") + f" {e}"

    fmt = _detect_format_ws(ws)
    if fmt == 'A':
        rows = _parse_format_a(ws, source_file)
    elif fmt == 'B':
        rows = _parse_format_b(ws, source_file)
    else:
        rows = _parse_format_c(ws, source_file)

    if not rows:
        return [], tr("Fajl prepoznat kao Format {a0}, ali nisu pronađeni podaci.", a0=fmt)

    rows = _apply_tax(rows, db)
    return rows, fmt


# ── routes ────────────────────────────────────────────────────────────

@bp.route("/")
def index():
    db = get_db()
    _ensure_indexes(db)

    filters = {
        "employee":    request.args.get("employee", ""),
        "venue":       request.args.get("venue", ""),
        "worker_type": request.args.get("worker_type", ""),
        "date_from":   request.args.get("date_from", ""),
        "date_to":     request.args.get("date_to", ""),
    }
    expand_month = request.args.get("expand", "")
    view_mode    = request.args.get("view_mode", "daily")

    monthly                   = _get_monthly_summary(db, filters)
    kpis                      = _kpis_from_monthly(monthly)   # no extra DB query
    employees, venues, worker_types = _get_filter_options(db) # 1 query instead of 3

    expanded_rows   = []
    aggregated_rows = []
    if expand_month:
        expanded_rows   = _get_month_detail(db, expand_month, filters)
        aggregated_rows = _aggregated_from_rows(expanded_rows)  # pure Python, no DB

    return render_template(
        "wages/index.html",
        tab="dnevnice",
        kpis=kpis,
        monthly=monthly,
        employees=employees,
        venues=venues,
        worker_types=worker_types,
        filters=filters,
        expand_month=expand_month,
        expanded_rows=expanded_rows,
        aggregated_rows=aggregated_rows,
        view_mode=view_mode,
    )


@bp.route("/upload", methods=["GET", "POST"])
def upload():
    db = get_db()

    if request.method == "GET":
        return render_template("wages/upload.html")

    f = request.files.get("file")
    if not f or not f.filename:
        flash(tr("Nije izabran fajl."), "error")
        return redirect(url_for("wages.upload"))

    # Save to temp location
    import tempfile
    suffix = os.path.splitext(f.filename)[1]
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    f.save(tmp.name)
    tmp.close()

    rows, fmt_or_error = _parse_file(tmp.name, f.filename, db)
    os.unlink(tmp.name)

    if not rows:
        flash(tr("Nije moguće parsirati fajl: {fmt_or_error}", fmt_or_error=fmt_or_error), "error")
        return redirect(url_for("wages.upload"))

    # Store rows in server-side session to avoid JSON encoding issues in forms
    from flask import session
    import json
    session["wages_preview"] = json.dumps(rows, ensure_ascii=False)
    session["wages_source"]  = f.filename

    return render_template(
        "wages/preview.html",
        rows=rows,
        source_file=f.filename,
        fmt=fmt_or_error,
        row_count=len(rows),
    )


@bp.route("/commit", methods=["POST"])
def commit():
    """Commit previewed rows to wages_daily."""
    import json
    from flask import session
    db = get_db()

    rows_json = session.pop("wages_preview", None)
    if not rows_json:
        flash(tr("Sesija je istekla. Molimo ponovo učitajte fajl."), "error")
        return redirect(url_for("wages.upload"))
    try:
        rows = json.loads(rows_json)
    except Exception:
        flash(tr("Greška pri čitanju podataka."), "error")
        return redirect(url_for("wages.upload"))

    inserted = 0
    skipped  = 0
    now      = _date.today().isoformat()

    for r in rows:
        try:
            db.execute("""
                INSERT OR IGNORE INTO wages_daily
                  (date, employee_name, worker_type, venue,
                   net_amount, tax_bracket, gross_cost,
                   source_file, notes, created_at, fitid)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)
            """, [
                r["date"], r["employee_name"], r.get("worker_type", ""),
                r.get("venue", ""), r["net_amount"],
                r.get("tax_bracket", "76%"), r.get("gross_cost", 0),
                r.get("source_file", ""), r.get("notes", ""),
                now, r.get("fitid", ""),
            ])
            if db.execute("SELECT changes()").fetchone()[0]:
                inserted += 1
            else:
                skipped += 1
        except Exception as e:
            flash(tr("Greška pri upisu: {e}", e=e), "error")
            db.rollback()
            return redirect(url_for("wages.upload"))

    db.commit()
    flash(tr("Upisano {inserted} redova. Preskočeno (duplikati): {skipped}.", inserted=inserted, skipped=skipped), "success")
    return redirect(url_for("wages.index"))


# ── Stalni zaposleni (salaries_monthly) ──────────────────────────────

def _ensure_salary_columns(db):
    """Add monthly_gross column if not present."""
    cols = [r[1] for r in db.execute("PRAGMA table_info(salaries_monthly)").fetchall()]
    if "monthly_gross" not in cols:
        db.execute("ALTER TABLE salaries_monthly ADD COLUMN monthly_gross REAL DEFAULT 0")
        db.commit()


def _calc_gross(net, db):
    """Compute gross from net using 76% bracket as default (conservative)."""
    try:
        return round(float(net) * 1.76, 2)
    except (TypeError, ValueError):
        return 0.0


def _get_salaries_with_history(db):
    """Get all employees grouped with their salary periods"""
    rows = db.execute("""
        SELECT 
            id, employee_name, role,
            monthly_amount AS monthly_net,
            monthly_gross,
            valid_from,
            valid_to,
            CASE WHEN valid_to IS NULL THEN 1 ELSE 0 END as is_active
        FROM salaries_monthly
        ORDER BY employee_name, valid_from DESC
    """).fetchall()
    
    # Group by employee
    from collections import defaultdict
    employees = defaultdict(lambda: {'periods': [], 'role': None})
    
    for r in rows:
        emp = employees[r['employee_name']]
        if emp['role'] is None:
            emp['role'] = r['role']
        emp['periods'].append(dict(r))
    
    # Convert to list with employee_name as field
    result = []
    for name, data in employees.items():
        result.append({
            'employee_name': name,
            'role': data['role'],
            'periods': data['periods'],
            'current_net': next((p['monthly_net'] for p in data['periods'] if p['is_active']), 0),
            'current_gross': next((p['monthly_gross'] for p in data['periods'] if p['is_active']), 0)
        })
    
    return sorted(result, key=lambda x: x['employee_name'])


@bp.route("/stalni")
def stalni():
    db = get_db()
    _ensure_salary_columns(db)
    
    employees = _get_salaries_with_history(db)
    
    total_net = sum(e['current_net'] for e in employees)
    total_gross = sum(e['current_gross'] for e in employees)
    
    return render_template(
        "wages/index.html",
        tab="stalni",
        employees=employees,
        total_net=total_net,
        total_gross=total_gross,
        kpis={}, monthly=[], employees_list=[], venues=[],
        worker_types=[], filters={}, expand_month="", expanded_rows=[],
        aggregated_rows=[], view_mode="daily"
    )


@bp.route("/stalni/add", methods=["POST"])
def stalni_add():
    db = get_db()
    _ensure_salary_columns(db)
    name  = request.form.get("employee_name", "").strip()
    role  = request.form.get("role", "").strip()
    net   = request.form.get("monthly_net", "0")
    gross = request.form.get("monthly_gross", "")
    valid_from = request.form.get("valid_from", "").strip() or _date.today().isoformat()

    if not name:
        flash(tr("Ime zaposlenog je obavezno."), "error")
        return redirect(url_for("wages.stalni"))
    try:
        net_val   = float(net)
        gross_val = float(gross) if gross else _calc_gross(net_val, db)
    except ValueError:
        flash(tr("Neispravan iznos."), "error")
        return redirect(url_for("wages.stalni"))

    db.execute("""
        INSERT INTO salaries_monthly (employee_name, role, monthly_amount, monthly_gross, valid_from, valid_to)
        VALUES (?, ?, ?, ?, ?, NULL)
    """, [name, role, net_val, gross_val, valid_from])
    db.commit()
    flash(tr("Dodat zaposleni: {name}.", name=name), "success")
    return redirect(url_for("wages.stalni"))




@bp.route("/stalni/add-period", methods=["POST"])
def stalni_add_period():
    """Add a new salary period for existing employee"""
    db = get_db()
    _ensure_salary_columns(db)
    
    employee_name = request.form.get("employee_name", "").strip()
    new_net = request.form.get("monthly_net", "0")
    new_gross = request.form.get("monthly_gross", "")
    valid_from = request.form.get("valid_from", "").strip()
    
    if not employee_name or not valid_from:
        flash(tr("Ime zaposlenog i datum početka su obavezni."), "error")
        return redirect(url_for("wages.stalni"))
    
    try:
        net_val = float(new_net)
        gross_val = float(new_gross) if new_gross else _calc_gross(net_val, db)
    except ValueError:
        flash(tr("Neispravan iznos."), "error")
        return redirect(url_for("wages.stalni"))
    
    # Close previous active period (set valid_to to day before new period)
    from datetime import datetime, timedelta
    new_start = datetime.fromisoformat(valid_from)
    prev_end = (new_start - timedelta(days=1)).strftime('%Y-%m-%d')
    
    db.execute("""
        UPDATE salaries_monthly
        SET valid_to = ?
        WHERE employee_name = ? AND valid_to IS NULL
    """, [prev_end, employee_name])
    
    # Get role from existing records
    role_row = db.execute(
        "SELECT role FROM salaries_monthly WHERE employee_name = ? LIMIT 1",
        [employee_name]
    ).fetchone()
    role = role_row['role'] if role_row else ""
    
    # Insert new period
    db.execute("""
        INSERT INTO salaries_monthly 
        (employee_name, role, monthly_amount, monthly_gross, valid_from, valid_to)
        VALUES (?, ?, ?, ?, ?, NULL)
    """, [employee_name, role, net_val, gross_val, valid_from])
    
    db.commit()
    flash(tr("Nova plata za {employee_name} od {valid_from}: {a2} RSD.", employee_name=employee_name, valid_from=valid_from, a2=_money(net_val)), "success")
    return redirect(url_for("wages.stalni"))

@bp.route("/stalni/<int:salary_id>/edit", methods=["POST"])
def stalni_edit(salary_id):
    db = get_db()
    _ensure_salary_columns(db)
    name  = request.form.get("employee_name", "").strip()
    role  = request.form.get("role", "").strip()
    net   = request.form.get("monthly_net", "0")
    gross = request.form.get("monthly_gross", "")

    if not name:
        flash(tr("Ime zaposlenog je obavezno."), "error")
        return redirect(url_for("wages.stalni"))
    try:
        net_val   = float(net)
        gross_val = float(gross) if gross else _calc_gross(net_val, db)
    except ValueError:
        flash(tr("Neispravan iznos."), "error")
        return redirect(url_for("wages.stalni"))

    db.execute("""
        UPDATE salaries_monthly
        SET employee_name=?, role=?, monthly_amount=?, monthly_gross=?
        WHERE id=?
    """, [name, role, net_val, gross_val, salary_id])
    db.commit()
    flash(tr("Ažurirano: {name}.", name=name), "success")
    return redirect(url_for("wages.stalni"))


@bp.route("/stalni/<int:salary_id>/delete", methods=["POST"])
def stalni_delete(salary_id):
    db = get_db()
    row = db.execute("SELECT employee_name FROM salaries_monthly WHERE id=?", [salary_id]).fetchone()
    if row:
        db.execute("DELETE FROM salaries_monthly WHERE id=?", [salary_id])
        db.commit()
        flash(tr("Obrisan: {a0}.", a0=row['employee_name']), "success")
    return redirect(url_for("wages.stalni"))


# ══════════════════════════════════════════════════════════════════════
# ZAPOSLENI (Employees) — Tax bracket management
# ══════════════════════════════════════════════════════════════════════
 
def _ensure_bracket_tables(db):
    """Create employee_bracket_history table if not exists."""
    db.execute("""
        CREATE TABLE IF NOT EXISTS employee_bracket_history (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_name TEXT    NOT NULL,
            bracket       TEXT    NOT NULL,
            valid_from    TEXT    NOT NULL,
            notes         TEXT,
            created_at    TEXT    DEFAULT (datetime('now')),
            UNIQUE(employee_name, valid_from)
        )
    """)
    db.commit()
 
 
def _gross_cost(net, bracket_pct_str):
    """Compute gross cost from net amount and bracket percentage string."""
    try:
        pct = float(str(bracket_pct_str).replace('%','').strip())
        return round(float(net) * (1 + pct / 100), 2)
    except (TypeError, ValueError):
        return 0.0
 
 
def _get_employees_with_brackets(db):
    """Get employees from wages_daily with their current bracket."""
    wage_rows = db.execute("""
        SELECT employee_name, worker_type,
               COUNT(id) as wage_rows,
               SUM(net_amount) as total_net,
               SUM(gross_cost) as total_gross
        FROM wages_daily
        GROUP BY employee_name, worker_type
        ORDER BY employee_name
    """).fetchall()
 
    bracket_rows = db.execute("""
        SELECT employee_name, bracket, valid_from
        FROM employee_bracket_history
        ORDER BY employee_name, valid_from DESC
    """).fetchall()
    
    brackets = {}
    for r in bracket_rows:
        name = r['employee_name']
        if name not in brackets:
            brackets[name] = dict(r)
 
    result = []
    for w in wage_rows:
        d = dict(w)
        b = brackets.get(d['employee_name'], {})
        d['bracket']    = b.get('bracket')
        d['valid_from'] = b.get('valid_from')
        result.append(d)
    return result
 
 
def _get_bracket_history(db):
    """Get all bracket history grouped by employee."""
    employees = _get_employees_with_brackets(db)
    history = {}
    for e in employees:
        name = e['employee_name']
        rows = db.execute("""
            SELECT * FROM employee_bracket_history
            WHERE employee_name = ?
            ORDER BY valid_from DESC
        """, [name]).fetchall()
        history[name] = [dict(r) for r in rows]
    return history
 
 
@bp.route("/zaposleni")
def zaposleni():
    """Employees tab — tax bracket management."""
    from datetime import date
    db = get_db()
    _ensure_bracket_tables(db)
    
    employees_data = _get_employees_with_brackets(db)
    history = _get_bracket_history(db)
    
    return render_template(
        "wages/index.html",
        tab="zaposleni",
        employees_data=employees_data,
        history=history,
        today=date.today().isoformat(),
        # Empty defaults for other tabs
        kpis={}, monthly=[], employees=[], venues=[],
        worker_types=[], filters={}, expand_month="", expanded_rows=[],
        aggregated_rows=[], view_mode="daily",
        total_net=0, total_gross=0
    )
 
 
@bp.route("/zaposleni/set-bracket", methods=["POST"])
def zaposleni_set_bracket():
    """Set or update employee tax bracket."""
    db = get_db()
    _ensure_bracket_tables(db)
    
    f             = request.form
    employee_name = f.get("employee_name","").strip()
    bracket_raw   = f.get("bracket","").strip().replace('%','')
    valid_from    = f.get("valid_from","")
    notes         = f.get("notes","").strip() or None
 
    try:
        bracket_pct = float(bracket_raw)
        if bracket_pct <= 0 or bracket_pct > 200:
            raise ValueError
    except ValueError:
        flash(tr("Unesite ispravnu stopu (broj između 1 i 200)."), "error")
        return redirect(url_for("wages.zaposleni"))
 
    if not employee_name or not valid_from:
        flash(tr("Popunite sva polja."), "error")
        return redirect(url_for("wages.zaposleni"))
 
    bracket_str = str(int(bracket_pct)) if bracket_pct == int(bracket_pct) else str(bracket_pct)
 
    try:
        db.execute("""
            INSERT INTO employee_bracket_history
                (employee_name, bracket, valid_from, notes)
            VALUES (?,?,?,?)
            ON CONFLICT(employee_name, valid_from) DO UPDATE SET
                bracket=excluded.bracket, notes=excluded.notes
        """, [employee_name, bracket_str, valid_from, notes])
        db.commit()
        flash(tr("Stopa {bracket_str}% za '{employee_name}' važi od {valid_from}.", bracket_str=bracket_str, employee_name=employee_name, valid_from=valid_from), "success")
    except Exception as e:
        flash(tr("Greška: {e}", e=e), "error")
        return redirect(url_for("wages.zaposleni"))
 
    # Recalculate gross_cost for this employee from valid_from onwards
    wages = db.execute("""
        SELECT id, date, net_amount FROM wages_daily
        WHERE employee_name = ? AND date >= ?
    """, [employee_name, valid_from]).fetchall()
 
    updated = 0
    for w in wages:
        b = _get_bracket_for_date(db, employee_name, w['date'])
        if b:
            gc = _gross_cost(w['net_amount'], b)
            db.execute("UPDATE wages_daily SET gross_cost=?, tax_bracket=? WHERE id=?",
                       [gc, b, w['id']])
            updated += 1
    db.commit()
    flash(tr("Preračunato {updated} stavki zarada.", updated=updated), "success")
    return redirect(url_for("wages.zaposleni"))
 
 
@bp.route("/zaposleni/recalc-all", methods=["POST"])
def zaposleni_recalc_all():
    """Recalculate all wage gross costs based on current brackets."""
    db = get_db()
    _ensure_bracket_tables(db)
    
    wages = db.execute("SELECT id, date, employee_name, net_amount FROM wages_daily").fetchall()
    updated = skipped = 0
    
    for w in wages:
        b = _get_bracket_for_date(db, w['employee_name'], w['date'])
        if b:
            gc = _gross_cost(w['net_amount'], b)
            db.execute("UPDATE wages_daily SET gross_cost=?, tax_bracket=? WHERE id=?",
                       [gc, b, w['id']])
            updated += 1
        else:
            skipped += 1
    
    db.commit()
    flash(tr("Preračunato {updated} stavki | {skipped} bez definisane stope.", updated=updated, skipped=skipped), "success")
    return redirect(url_for("wages.zaposleni"))
 
 
@bp.route("/zaposleni/delete-bracket/<int:bracket_id>", methods=["POST"])
def zaposleni_delete_bracket(bracket_id):
    """Delete a bracket history entry."""
    db = get_db()
    db.execute("DELETE FROM employee_bracket_history WHERE id=?", [bracket_id])
    db.commit()
    flash(tr("Stopa obrisana."), "success")
    return redirect(url_for("wages.zaposleni"))